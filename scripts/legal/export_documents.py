"""Export the app's legal JSON into editable Word, PDF and Markdown documents.
Requires python-docx and reportlab. No network or live customer data is used.
"""
import json
from html import escape
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'docs' / 'legal'
BLUE = '#2458B8'
INK = '#17243A'
MUTED = '#526078'
font_dir = Path('/System/Library/Fonts/Supplemental')
if (font_dir / 'Arial.ttf').exists():
    pdfmetrics.registerFont(TTFont('LegalBody', str(font_dir / 'Arial.ttf')))
    pdfmetrics.registerFont(TTFont('LegalBold', str(font_dir / 'Arial Bold.ttf')))
    pdfmetrics.registerFontFamily('LegalBody', normal='LegalBody', bold='LegalBold')
    regular, bold = 'LegalBody', 'LegalBold'
else:
    regular, bold = 'Helvetica', 'Helvetica-Bold'

styles = {
    'brand': ParagraphStyle('brand', fontName=bold, fontSize=11, leading=15, textColor=colors.HexColor(BLUE), spaceAfter=10),
    'title': ParagraphStyle('title', fontName=bold, fontSize=29, leading=34, textColor=colors.HexColor(INK), spaceAfter=12),
    'meta': ParagraphStyle('meta', fontName=regular, fontSize=9, leading=14, textColor=colors.HexColor(MUTED), spaceAfter=12),
    'notice': ParagraphStyle('notice', fontName=regular, fontSize=9, leading=13, textColor=colors.HexColor(INK), backColor=colors.HexColor('#EFF4FC'), borderPadding=10, spaceBefore=5, spaceAfter=20),
    'heading': ParagraphStyle('heading', fontName=bold, fontSize=12, leading=16, textColor=colors.HexColor(INK), spaceBefore=12, spaceAfter=7, keepWithNext=True),
    'body': ParagraphStyle('body', fontName=regular, fontSize=10, leading=14.5, textColor=colors.HexColor(INK), spaceAfter=9, allowWidows=0, allowOrphans=0, alignment=TA_LEFT),
}

def footer(canvas, doc):
    canvas.saveState()
    width, height = A4
    canvas.setFillColor(colors.white)
    canvas.rect(0, 0, width, height, fill=1, stroke=0)
    canvas.setStrokeColor(colors.HexColor('#DCE3EE'))
    canvas.line(48, 43, width - 48, 43)
    canvas.setFillColor(colors.HexColor(MUTED))
    canvas.setFont(regular, 8)
    canvas.drawString(48, 29, f'LifeSwap | {doc.title} | Review draft')
    canvas.drawRightString(width - 48, 29, str(doc.page))
    canvas.restoreState()

def export_pdf(data, path):
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=48, leftMargin=48,
        topMargin=44, bottomMargin=58, title=data['title'], author='Itay Simchai',
        subject='LifeSwap legal document - review draft')
    story = [Paragraph('LIFESWAP', styles['brand']), Paragraph(data['title'], styles['title']),
        Paragraph(f"Review draft | Updated {escape(data['updated'])}<br/>Itay Simchai, Israel | nadrty8@gmail.com<br/>Postal address: pending | Effective date: pending publication", styles['meta']),
        Paragraph(escape(data['status']), styles['notice'])]
    for section in data['sections']:
        story.append(Paragraph(escape(section['heading']), styles['heading']))
        for paragraph in section['paragraphs']:
            story.append(KeepTogether([Paragraph(escape(paragraph), styles['body'])]))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)

def export_word(data, path):
    doc = Document()
    doc.core_properties.title = data['title']
    doc.core_properties.author = 'Itay Simchai'
    doc.core_properties.subject = 'LifeSwap legal document - review draft'
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.7)
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(INK[1:])
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(9)
    normal.paragraph_format.widow_control = True
    normal.paragraph_format.keep_together = True
    heading = doc.styles['Heading 1']
    heading.font.name = 'Arial'
    heading.font.size = Pt(12)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor.from_string(INK[1:])
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(7)
    heading.paragraph_format.keep_with_next = True
    brand = doc.add_paragraph()
    run = brand.add_run('LIFESWAP')
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE[1:])
    title = doc.add_paragraph()
    run = title.add_run(data['title'])
    run.bold = True
    run.font.size = Pt(29)
    meta = doc.add_paragraph(f"Review draft | Updated {data['updated']}\nItay Simchai, Israel | nadrty8@gmail.com\nPostal address: pending | Effective date: pending publication")
    for run in meta.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(MUTED[1:])
    notice = doc.add_paragraph(data['status'])
    shade = OxmlElement('w:shd')
    shade.set(qn('w:fill'), 'EFF4FC')
    notice._p.get_or_add_pPr().append(shade)
    for run in notice.runs:
        run.font.size = Pt(9)
    for item in data['sections']:
        doc.add_heading(item['heading'], level=1)
        for paragraph in item['paragraphs']:
            doc.add_paragraph(paragraph)
    foot = section.footer.paragraphs[0]
    foot.text = f"LifeSwap | {data['title']} | Review draft   -   "
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    foot._p.append(field)
    for run in foot.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED[1:])
    doc.save(path)

for slug, filename in [('privacy-policy', 'LifeSwap-Privacy-Policy'), ('terms-of-use', 'LifeSwap-Terms-of-Use')]:
    data = json.loads((ROOT / 'src/content/legal' / f'{slug}.json').read_text())
    export_pdf(data, OUT / f'{filename}.pdf')
    export_word(data, OUT / f'{filename}.docx')
    markdown = f"# LifeSwap {data['title']}\n\nDraft updated: {data['updated']}\n\n{data['status']}\n\n"
    markdown += '\n\n'.join('## ' + s['heading'] + '\n\n' + '\n\n'.join(s['paragraphs']) for s in data['sections']) + '\n'
    (OUT / f'{slug}.md').write_text(markdown)
    print(f'Exported {filename}.pdf and {filename}.docx')
